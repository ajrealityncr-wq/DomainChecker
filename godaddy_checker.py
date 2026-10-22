"""Check domain availability using the official GoDaddy Domains API.

Settings come from settings.py. The only inputs are the domain list (--input) and the
GoDaddy token (GODADDY_PAT in the environment or .env).
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable

import requests

import settings
from history import DomainHistory

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - requirements.txt installs this package
    load_dotenv = None


DOMAIN_RE = re.compile(
    r"^(?=.{1,253}$)(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}$",
    re.IGNORECASE,
)
RETRYABLE_STATUS_CODES = {408, 429, 500, 502, 503, 504}
AUTH_STATUS_CODES = {401, 403}
RDAP_REGISTERED = "RDAP_REGISTERED"
PREVIOUSLY_CHECKED = "PREVIOUSLY_CHECKED"
logger = logging.getLogger("godaddy_checker")
RECORD_FIELDS = {"domain", "available", "price", "renewal_price", "currency", "premium", "definitive", "status", "restrictions", "checked_at", "error", "api_success"}
CSV_FIELDS = ["domain", "available", "price", "renewal_price", "currency", "premium", "definitive", "status", "checked_at"]
EXCEL_HEADERS = ["input_domain", "domain", "status", "reason", "available", "price", "renewal_price", "currency", "premium", "definitive", "checked_at", "error"]
EXCEL_WIDTHS = [24, 30, 14, 30, 12, 12, 14, 10, 12, 12, 22, 50]
EXCEL_STATUS_FILLS = {"AVAILABLE": "C6EFCE", "TAKEN": "FFC7CE", "INVALID": "FFEB9C", "ERROR": "F4CCCC", "UNCHECKED": "D9EAF7"}
MICRO_UNIT_THRESHOLD = 1_000_000


class AuthError(RuntimeError):
    """GoDaddy rejected the token. Continuing would only waste requests, so the run stops."""


@dataclass
class CheckResult:
    domain: str
    available: bool | None
    price: Any = None
    renewal_price: Any = None
    currency: str | None = None
    premium: Any = None
    definitive: Any = None
    status: str | None = None
    restrictions: Any = None
    checked_at: str = ""
    error: str | None = None
    api_success: bool = False
    raw: dict[str, Any] = field(default_factory=dict, repr=False)


def default_config() -> dict[str, Any]:
    """The checker's settings, built from settings.py. Nothing here is read from a file."""
    return {
        "batch_size": settings.BATCH_SIZE,
        "optimize_for": settings.OPTIMIZE_FOR,
        "timeout_seconds": settings.TIMEOUT_SECONDS,
        "delay_between_batches_seconds": settings.DELAY_BETWEEN_BATCHES_SECONDS,
        "max_retries": settings.MAX_RETRIES,
        "retry_backoff_seconds": settings.RETRY_BACKOFF_SECONDS,
        "api_url": settings.GODADDY_API_URL,
        "output_directory": str(settings.OUTPUT_DIRECTORY),
        "checkpoint_file": str(settings.CHECKPOINT_FILE),
        "allowed_tlds": list(settings.ALLOWED_TLDS),
        "sort_available_by": "domain",
        "budget": {
            "max_yearly_price": settings.DEFAULT_MAX_PRICE,
            "max_renewal_price": None,
            "currency": settings.CURRENCY,
            "require_renewal_price": False,
        },
        "prefilter": {
            "rdap_url": settings.RDAP_URL,
            "timeout_seconds": settings.RDAP_TIMEOUT_SECONDS,
        },
    }


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def allowed_tld_set(config: dict[str, Any]) -> set[str]:
    return {"." + str(tld).lower().lstrip(".") for tld in config.get("allowed_tlds") or []}


def normalize_domain(value: Any, config: dict[str, Any]) -> tuple[str | None, str | None]:
    if not isinstance(value, str):
        return None, str(value)
    original = value
    domain = re.sub(r"^https?://", "", value.strip().lower()).split("/", 1)[0]
    if not domain or not DOMAIN_RE.fullmatch(domain):
        return None, original
    allowed = allowed_tld_set(config)
    if allowed and "." + domain.rsplit(".", 1)[1] not in allowed:
        return None, original
    return domain, original if domain != original else None


def prepare_domains(values: Iterable[Any], config: dict[str, Any]) -> tuple[list[str], list[str], list[dict[str, str]]]:
    valid: list[str] = []
    invalid: list[str] = []
    normalization_log: list[dict[str, str]] = []
    seen: set[str] = set()
    for value in values:
        domain, original = normalize_domain(value, config)
        if domain is None:
            invalid.append(original or "")
            continue
        if original is not None:
            normalization_log.append({"original": original, "normalized": domain})
        if domain in seen:
            continue
        seen.add(domain)
        valid.append(domain)
    return valid, invalid, normalization_log


def load_domains(path: Path) -> list[Any]:
    with path.open(encoding="utf-8") as handle:
        data = json.load(handle)
    if isinstance(data, list):
        if all(isinstance(item, dict) and isinstance(item.get("domain"), str) for item in data):
            return [item["domain"] for item in data]
        return data
    if isinstance(data, dict) and isinstance(data.get("domains"), list):
        return data["domains"]
    raise ValueError("Input JSON must be an array or an object containing a 'domains' array")


def batches(domains: list[str], size: int) -> Iterable[list[str]]:
    for index in range(0, len(domains), size):
        yield domains[index : index + size]


def pending_domains(valid: list[str], processed: dict[str, Any]) -> list[str]:
    """Domains that still need an answer. Errored domains are included so a resume retries them."""
    return [domain for domain in valid if domain not in processed]


def unchecked_domains(valid: list[str], processed: dict[str, Any], errors: list[dict[str, Any]]) -> list[str]:
    """Domains never attempted, for example because the run was stopped mid-way."""
    done = set(processed) | {item.get("domain") for item in errors}
    return [domain for domain in valid if domain not in done]


def record_errors(errors: list[dict[str, Any]], batch: list[str], message: str) -> list[dict[str, Any]]:
    """Replace any earlier error for these domains so each domain has at most one error entry."""
    batch_set = set(batch)
    kept = [item for item in errors if item.get("domain") not in batch_set]
    return kept + [
        {"domain": domain, "status": "ERROR", "reason": "API_CHECK_FAILED", "error": message, "checked_at": utc_now()}
        for domain in batch
    ]


def drop_resolved_errors(errors: list[dict[str, Any]], processed: dict[str, Any]) -> list[dict[str, Any]]:
    return [item for item in errors if item.get("domain") not in processed]


def load_pat() -> str | None:
    """The token comes from the environment or the project's .env file. It is never written to files or logs.

    .env is read here, not only in the command-line entry point, so the web app finds the token too.
    Variables already set in the environment take precedence over .env.
    """
    if load_dotenv:
        load_dotenv(settings.ROOT / ".env")
    token = os.getenv("GODADDY_PAT", "").strip()
    return token or None


def backoff_delay(config: dict[str, Any], attempt: int) -> float:
    return min(float(settings.MAX_RETRY_DELAY_SECONDS), float(config["retry_backoff_seconds"]) * (2**attempt))


def retry_delay(response: requests.Response, config: dict[str, Any], attempt: int) -> float:
    """Honour a numeric Retry-After header (capped); otherwise fall back to exponential backoff."""
    try:
        return min(float(settings.MAX_RETRY_DELAY_SECONDS), max(0.0, float(response.headers.get("Retry-After", ""))))
    except ValueError:
        return backoff_delay(config, attempt)


def check_batch(batch: list[str], config: dict[str, Any], pat: str, logger: logging.Logger) -> tuple[list[CheckResult], str | None]:
    """Return (results, error). Raises AuthError when the token is rejected."""
    payload = {"domains": batch, "optimizeFor": config["optimize_for"]}
    headers = {"Authorization": f"Bearer {pat}", "Content-Type": "application/json", "Accept": "application/json"}
    retries = int(config["max_retries"])
    for attempt in range(retries + 1):
        try:
            response = requests.post(config["api_url"], headers=headers, json=payload, timeout=float(config["timeout_seconds"]))
        except requests.RequestException as exc:
            if attempt < retries:
                delay = backoff_delay(config, attempt)
                logger.warning("Temporary network error; retrying in %.1f seconds: %s", delay, exc)
                time.sleep(delay)
                continue
            return [], f"network error: {exc}"
        if response.status_code in RETRYABLE_STATUS_CODES and attempt < retries:
            delay = retry_delay(response, config, attempt)
            logger.warning("GoDaddy returned HTTP %s; retrying in %.1f seconds", response.status_code, delay)
            time.sleep(delay)
            continue
        if response.status_code in AUTH_STATUS_CODES:
            raise AuthError(f"authentication/authorization failed (HTTP {response.status_code}); check GODADDY_PAT and its permissions")
        if not response.ok:
            return [], f"GoDaddy returned HTTP {response.status_code}"
        try:
            body = response.json()
        except ValueError:
            return [], "GoDaddy returned invalid JSON"
        if isinstance(body, dict) and isinstance(body.get("items"), list):
            body = body["items"]
        if not isinstance(body, list):
            return [], "GoDaddy response must be an array or an object containing an items array"
        return parse_response(body, batch)
    return [], "request retries exhausted"


def parse_response(items: list[Any], requested: list[str]) -> tuple[list[CheckResult], str | None]:
    requested_set = set(requested)
    results: list[CheckResult] = []
    returned: set[str] = set()
    for item in items:
        if not isinstance(item, dict) or not isinstance(item.get("domain"), str):
            return [], "malformed response item"
        domain = item["domain"].strip().lower()
        if domain not in requested_set or domain in returned:
            return [], f"unexpected or duplicate domain in response: {domain}"
        returned.add(domain)
        available = item.get("available")
        if not isinstance(available, bool):
            return [], f"missing or invalid availability for {domain}"
        results.append(CheckResult(
            domain=domain,
            available=available,
            price=item.get("price"),
            renewal_price=item.get("renewalPrice", item.get("renewal_price")),
            currency=item.get("currency"),
            premium=item.get("premium"),
            definitive=item.get("definitive"),
            status=item.get("status"),
            restrictions=item.get("restrictions"),
            checked_at=utc_now(),
            api_success=True,
            raw=item,
        ))
    missing = requested_set - returned
    if missing:
        return [], f"response omitted domains: {', '.join(sorted(missing))}"
    return results, None


def price_units_warning(raw_item: dict[str, Any]) -> str | None:
    """Flag prices that look like micro-units (GoDaddy-style 1,000,000 per unit), which would break budget checks."""
    price = raw_item.get("price")
    if isinstance(price, (int, float)) and price >= MICRO_UNIT_THRESHOLD:
        return (
            f"Price {price} looks like micro-units (a whole-currency amount multiplied by 1,000,000). "
            "Budget comparisons would be wrong until the units are confirmed."
        )
    return None


def rdap_is_registered(domain: str, config: dict[str, Any]) -> bool | None:
    """True if the RDAP registry knows the domain, False if it reports not found, None if inconclusive.

    Only a True result is acted on; availability always comes from GoDaddy.
    """
    rdap = config["prefilter"]
    try:
        response = requests.get(
            rdap["rdap_url"].format(domain=domain),
            headers={"Accept": "application/rdap+json"},
            timeout=float(rdap["timeout_seconds"]),
        )
    except requests.RequestException:
        return None
    if response.status_code == 200:
        return True
    if response.status_code == 404:
        return False
    return None


def rdap_statuses(domains: list[str], config: dict[str, Any], max_workers: int = settings.RDAP_WORKERS,
                   on_progress: Callable[[int], None] | None = None) -> dict[str, bool | None]:
    """RDAP answer for each domain, looked up in parallel. A failed lookup counts as inconclusive (None)."""
    status: dict[str, bool | None] = {}
    if not domains:
        return status
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        futures = {pool.submit(rdap_is_registered, domain, config): domain for domain in domains}
        for done, future in enumerate(as_completed(futures), 1):
            domain = futures[future]
            try:
                status[domain] = future.result()
            except Exception:
                logger.exception("RDAP lookup failed for %s", domain)
                status[domain] = None
            if on_progress and (done % 10 == 0 or done == len(domains)):
                on_progress(done)
    return status


def check_in_batches(domains: list[str], config: dict[str, Any], pat: str, check_fn: Callable[..., Any], log: logging.Logger,
                     on_batch: Callable[[list[str], list["CheckResult"], str | None, int], None] | None = None,
                     ) -> tuple[list["CheckResult"], list[str]]:
    """Check domains in batches. Returns (results, failed domains).

    on_batch(batch, results, error, checked_so_far) runs after each batch, so callers can save progress.
    AuthError from check_fn is not caught here; it stops the whole run.
    """
    results: list[CheckResult] = []
    failed: list[str] = []
    checked = 0
    batch_list = list(batches(domains, config["batch_size"]))
    for index, batch in enumerate(batch_list, 1):
        batch_results, error = check_fn(batch, config, pat, log)
        if error:
            log.warning("Batch failed: %s", error)
            failed.extend(batch)
        else:
            results.extend(batch_results)
        checked += len(batch)
        if on_batch:
            on_batch(batch, batch_results, error, checked)
        if index < len(batch_list):
            time.sleep(float(config["delay_between_batches_seconds"]))
    return results, failed


def checkpoint_path(config: dict[str, Any], root: Path) -> Path:
    path = Path(str(config["checkpoint_file"]))
    return path if path.is_absolute() else root / path


def load_checkpoint(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"processed": {}, "api_requests": 0}
    try:
        with path.open(encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, dict) else {"processed": {}, "api_requests": 0}
    except (OSError, ValueError):
        return {"processed": {}, "api_requests": 0}


def save_checkpoint(path: Path, checkpoint: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(checkpoint, handle, indent=2)
    temporary.replace(path)


def record_for_output(result: CheckResult) -> dict[str, Any]:
    return {key: value for key, value in asdict(result).items() if key not in {"error", "api_success", "raw"} and value is not None}


def result_from_record(record: dict[str, Any]) -> CheckResult:
    return CheckResult(**{key: value for key, value in record.items() if key in RECORD_FIELDS})


def results_from_processed(processed: dict[str, dict[str, Any]]) -> list[CheckResult]:
    return [result_from_record(record) for record in processed.values() if record.get("available") in {True, False}]


def write_json(path: Path, data: Any) -> None:
    with path.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, indent=2)


def write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    fields = list(CSV_FIELDS)
    for record in records:
        for field_name in record:
            if field_name not in fields:
                fields.append(field_name)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            writer.writerow({field_name: record.get(field_name, "") for field_name in fields})


def markdown_value(value: Any) -> str:
    return str(value if value is not None and value != "" else "-").replace("|", "\\|").replace("\n", " ")


def write_markdown(path: Path, title: str, rows: list[dict[str, Any]]) -> None:
    headers = ["Domain", "Status", "Reason", "Available", "Price", "Renewal Price", "Currency", "Premium", "Definitive", "Checked At", "Error"]
    lines = [
        f"# {title}",
        "",
        f"Total: **{len(rows)}**",
        "",
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(["---"] * len(headers)) + "|",
    ]
    for row in rows:
        lines.append("| " + " | ".join([
            markdown_value(row.get("domain") or row.get("input_domain")),
            markdown_value(row.get("status")),
            markdown_value(row.get("reason")),
            markdown_value(row.get("available")),
            markdown_value(row.get("price")),
            markdown_value(row.get("renewal_price")),
            markdown_value(row.get("currency")),
            markdown_value(row.get("premium")),
            markdown_value(row.get("definitive")),
            markdown_value(row.get("checked_at")),
            markdown_value(row.get("error")),
        ]) + " |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_excel(path: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Domain Results"
    sheet.append(EXCEL_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([row.get(header, "") for header in EXCEL_HEADERS])
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(EXCEL_HEADERS))
    table = Table(displayName="DomainResultsTable", ref=f"A1:{last_column}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    for index, width in enumerate(EXCEL_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    status_index = EXCEL_HEADERS.index("status")
    for row in sheet.iter_rows(min_row=2):
        fill = EXCEL_STATUS_FILLS.get(row[status_index].value)
        if fill:
            row[status_index].fill = PatternFill("solid", fgColor=fill)
    workbook.save(path)


def has_limits(config: dict[str, Any]) -> bool:
    budget = config.get("budget") or {}
    return budget.get("max_yearly_price") is not None or budget.get("max_renewal_price") is not None


def meets_budget(record: dict[str, Any], config: dict[str, Any]) -> bool:
    """True when the record is under every cap that is set. With no caps set, every record passes.

    The first-year cap needs a known price in the required currency. The renewal cap is checked
    only when a renewal price is known, unless require_renewal_price is set.
    """
    budget = config.get("budget") or {}
    max_price = budget.get("max_yearly_price")
    max_renewal = budget.get("max_renewal_price")
    if max_price is None and max_renewal is None:
        return True
    if max_price is not None:
        price = record.get("price")
        currency = record.get("currency")
        required_currency = str(budget.get("currency", "")).upper()
        if not isinstance(price, (int, float)) or not currency or str(currency).upper() != required_currency:
            return False
        if price > float(max_price):
            return False
    if max_renewal is not None:
        renewal_price = record.get("renewal_price")
        if isinstance(renewal_price, (int, float)):
            if renewal_price > float(max_renewal):
                return False
        elif budget.get("require_renewal_price", False):
            return False
    return True


def status_for(result: CheckResult | None, error: dict[str, Any] | None) -> tuple[str, str]:
    if result and result.status == RDAP_REGISTERED:
        return "TAKEN", RDAP_REGISTERED
    if result and result.status == PREVIOUSLY_CHECKED:
        return ("AVAILABLE" if result.available else "TAKEN"), "CHECKED_EARLIER"
    if result and result.available is True:
        return "AVAILABLE", "GODADDY_CONFIRMED_AVAILABLE"
    if result and result.available is False:
        return "TAKEN", "GODADDY_CONFIRMED_UNAVAILABLE"
    if error:
        return "ERROR", error.get("reason", "API_CHECK_FAILED")
    return "UNCHECKED", "NOT_PROCESSED"


def excel_row(domain: str, result: CheckResult | None, error: dict[str, Any] | None) -> dict[str, Any]:
    status, reason = status_for(result, error)
    return {
        "input_domain": domain,
        "domain": domain,
        "status": status,
        "reason": reason,
        "available": result.available if result else "",
        "price": result.price if result else "",
        "renewal_price": result.renewal_price if result else "",
        "currency": result.currency if result else "",
        "premium": result.premium if result else "",
        "definitive": result.definitive if result else "",
        "checked_at": result.checked_at if result else (error or {}).get("checked_at", ""),
        "error": (error or {}).get("error", ""),
    }


def build_report_rows(valid: list[str], invalid: list[str], results: list[CheckResult], errors: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result_map = {item.domain: item for item in results}
    error_map = {item["domain"]: item for item in errors if item.get("domain")}
    rows = [excel_row(domain, result_map.get(domain), error_map.get(domain)) for domain in valid]
    rows.extend({
        "input_domain": value,
        "domain": "",
        "status": "INVALID",
        "reason": "INVALID_DOMAIN_SYNTAX",
        "available": False,
        "error": "Domain is not a valid domain name and was not sent to GoDaddy.",
    } for value in invalid)
    return rows


def sort_available(records: list[dict[str, Any]], mode: str) -> None:
    if mode == "domain":
        records.sort(key=lambda item: item["domain"])
    elif mode == "price":
        records.sort(key=lambda item: (item.get("price") is None, item.get("price") or 0))


def save_outputs(root: Path, config: dict[str, Any], valid: list[str], invalid: list[str], results: list[CheckResult], errors: list[dict[str, Any]], unchecked: list[str], started_at: str, requests_made: int, normalization_log: list[dict[str, str]]) -> dict[str, Any]:
    output = root / str(config["output_directory"])
    output.mkdir(parents=True, exist_ok=True)
    available = [record_for_output(item) for item in results if item.available is True]
    taken = [record_for_output(item) for item in results if item.available is False]
    sort_available(available, config["sort_available_by"])
    invalid = list(dict.fromkeys(invalid))
    write_json(output / "available_domains.json", available)
    write_json(output / "taken_domains.json", taken)
    write_json(output / "invalid_domains.json", invalid)
    write_json(output / "error_domains.json", errors)
    write_json(output / "unchecked_domains.json", unchecked)
    write_json(output / "normalization_log.json", normalization_log)
    write_csv(output / "available_domains.csv", available)
    write_csv(output / "taken_domains.csv", taken)
    write_csv(output / "invalid_domains.csv", [{"domain": value} for value in invalid])
    write_csv(output / "error_domains.csv", errors)
    write_csv(output / "unchecked_domains.csv", [{"domain": value} for value in unchecked])
    rows = build_report_rows(valid, invalid, results, errors)
    available_rows = [row for row in rows if row["status"] == "AVAILABLE"]
    budget_rows = [row for row in available_rows if meets_budget(row, config)]
    write_markdown(output / "domain_results.md", "Domain Check Results", rows)
    write_markdown(output / "available_domains.md", "Available Domains", available_rows)
    write_markdown(output / "budget_available_domains.md", "Budget-Qualified Available Domains", budget_rows)
    write_excel(output / "domain_results.xlsx", rows)
    write_excel(output / "available_domains.xlsx", available_rows)
    write_excel(output / "budget_available_domains.xlsx", budget_rows)
    buy_now = [item for item in available if meets_budget(item, config)]
    with (output / "BUY_NOW.txt").open("w", encoding="utf-8") as handle:
        handle.write("\n".join(item["domain"] for item in buy_now))
        if buy_now:
            handle.write("\n")
    summary = {
        "total_input": len(valid) + len(invalid),
        "unique_domains": len(valid),
        "invalid_domains": len(invalid),
        "processed": len(results),
        "available": len(available),
        "taken": len(taken),
        "rdap_registered": sum(1 for item in results if item.status == RDAP_REGISTERED),
        "errors": len(errors),
        "unchecked": len(unchecked),
        "api_requests": requests_made,
        "started_at": started_at,
        "finished_at": utc_now(),
    }
    write_json(output / "summary.json", summary)
    return summary


def configure_logging(root: Path, config: dict[str, Any]) -> logging.Logger:
    logger = logging.getLogger("godaddy_checker")
    logger.handlers.clear()
    logger.setLevel(logging.INFO)
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    console = logging.StreamHandler()
    console.setFormatter(formatter)
    logger.addHandler(console)
    output = root / str(config["output_directory"])
    output.mkdir(parents=True, exist_ok=True)
    file_handler = logging.FileHandler(output / "checker.log", encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    return logger


def print_summary(output: Path, summary: dict[str, Any]) -> None:
    print("\n========================================")
    print("GO DADDY DOMAIN CHECK COMPLETE")
    print("========================================")
    rows = [
        ("Input domains", "total_input"),
        ("Unique valid", "unique_domains"),
        ("Invalid", "invalid_domains"),
        ("Checked", "processed"),
        ("Available", "available"),
        ("Taken", "taken"),
        ("RDAP registered", "rdap_registered"),
        ("Errors", "errors"),
        ("Unchecked", "unchecked"),
        ("API requests", "api_requests"),
    ]
    for label, key in rows:
        print(f"{label:<20}: {summary[key]}")
    print(f"\nAvailable domains saved to:\n{output / 'available_domains.json'}")
    print(f"\nBUY NOW list:\n{output / 'BUY_NOW.txt'}")
    print("========================================")
    print("IMPORTANT:")
    print("GoDaddy availability results are for search purposes. Re-verify availability and price at checkout.")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Check domain availability with the official GoDaddy API")
    parser.add_argument("--input", help=f"JSON list of domains (default: {settings.INPUT_FILE})")
    parser.add_argument("--dry-run", action="store_true", help="Validate the list only; no network request")
    parser.add_argument("--test", action="store_true", help="Check one domain and print the raw API item")
    parser.add_argument("--force-recheck", action="store_true",
                        help="Ignore this run's saved checkpoint (names GoDaddy already answered are still not re-sent)")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    root = settings.ROOT
    if load_dotenv:
        load_dotenv(root / ".env")
    config = default_config()
    try:
        logger = configure_logging(root, config)
        input_path = Path(args.input) if args.input else root / settings.INPUT_FILE
        values = load_domains(input_path)
        valid, invalid, normalization_log = prepare_domains(values, config)
    except (OSError, ValueError) as exc:
        print(f"Configuration/input error: {exc}", file=sys.stderr)
        return 2

    output_dir = root / str(config["output_directory"])
    expected = (len(valid) + config["batch_size"] - 1) // config["batch_size"]
    print(f"Valid unique: {len(valid)} | Invalid: {len(invalid)} | Batch size: {config['batch_size']}")
    print(f"Expected API requests if nothing is resumed: {expected}")
    if args.dry_run:
        save_outputs(root, config, valid, invalid, [], [], valid, utc_now(), 0, normalization_log)
        print("Dry run complete. No GoDaddy or RDAP request was made.")
        print(f"Markdown report saved to: {output_dir / 'domain_results.md'}")
        return 0

    pat = load_pat()
    if not pat:
        missing_pat_errors = [
            {
                "domain": domain,
                "status": "ERROR",
                "reason": "MISSING_GODADDY_PAT",
                "error": "GoDaddy PAT is not configured; API check was not attempted.",
                "checked_at": utc_now(),
            }
            for domain in valid
        ]
        save_outputs(root, config, valid, invalid, [], missing_pat_errors, [], utc_now(), 0, normalization_log)
        print("Missing GoDaddy PAT. Set GODADDY_PAT in the environment or in .env.", file=sys.stderr)
        print(f"Results with reasons saved to: {output_dir}", file=sys.stderr)
        return 2
    if args.test:
        if not valid:
            print("Test failed: no valid domain is available in the input file.", file=sys.stderr)
            return 2
        try:
            results, error = check_batch(valid[:1], config, pat, logger)
        except AuthError as exc:
            print(f"API test failed: {exc}", file=sys.stderr)
            return 1
        if error:
            print(f"API test failed: {error}", file=sys.stderr)
            return 1
        print(f"API test succeeded for {results[0].domain}.")
        print("Raw GoDaddy item (confirm price units and currency before trusting budget filters):")
        print(json.dumps(results[0].raw, indent=2))
        warning = price_units_warning(results[0].raw)
        if warning:
            print(f"WARNING: {warning}", file=sys.stderr)
            return 1
        return 0

    started_at = utc_now()
    checkpoint_file = checkpoint_path(config, root)
    checkpoint = {"processed": {}, "api_requests": 0} if args.force_recheck else load_checkpoint(checkpoint_file)
    processed: dict[str, dict[str, Any]] = checkpoint.setdefault("processed", {})
    errors: list[dict[str, Any]] = list(checkpoint.get("errors", []))
    requests_made = int(checkpoint.get("api_requests", 0))
    results = results_from_processed(processed)

    def save_progress() -> None:
        checkpoint.update({"processed": processed, "api_requests": requests_made, "errors": errors})
        save_checkpoint(checkpoint_file, checkpoint)

    def write_reports() -> dict[str, Any]:
        return save_outputs(
            root, config, valid, invalid, results, errors,
            unchecked_domains(valid, processed, errors),
            started_at, requests_made, normalization_log,
        )

    history = DomainHistory()
    # Names answered in any earlier run or web search are never sent to GoDaddy again.
    earlier = history.checked(pending_domains(valid, processed))
    for domain, available in earlier.items():
        processed[domain] = record_for_output(CheckResult(domain=domain, available=available, status=PREVIOUSLY_CHECKED))
    if earlier:
        logger.info("Skipping %s domains already checked earlier", len(earlier))

    candidates = pending_domains(valid, processed)
    logger.info("RDAP pre-check for %s domains", len(candidates))
    registered_now = {}
    for domain, registered in rdap_statuses(candidates, config).items():
        if registered is True:
            logger.info("Taken (RDAP): %s", domain)
            processed[domain] = record_for_output(CheckResult(domain=domain, available=False, status=RDAP_REGISTERED, checked_at=utc_now()))
            registered_now[domain] = False
    history.mark_checked(registered_now)
    results = results_from_processed(processed)
    save_progress()

    pending = pending_domains(valid, processed)
    logger.info("Pending: %s domains in %s batches", len(pending), -(-len(pending) // config["batch_size"]))

    def on_batch(batch: list[str], batch_results: list[CheckResult], error: str | None, _checked: int) -> None:
        nonlocal errors, results, requests_made
        requests_made += 1
        if error:
            errors = record_errors(errors, batch, error)
        else:
            for result in batch_results:
                processed[result.domain] = record_for_output(result) | {"api_success": True}
                logger.info("%s: %s", "Available" if result.available else "Taken", result.domain)
            history.mark_checked({result.domain: result.available is True for result in batch_results})
            errors = drop_resolved_errors(errors, processed)
            results = results_from_processed(processed)
        save_progress()

    try:
        check_in_batches(pending, config, pat, check_batch, logger, on_batch)
    except AuthError as exc:
        logger.error("Run stopped: %s", exc)
        save_progress()
        write_reports()
        print(f"Stopped: {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        save_progress()
        write_reports()
        print("\nProcess interrupted. Progress has been saved; run again to resume.")
        return 130

    save_progress()
    summary = write_reports()
    print_summary(output_dir, summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
